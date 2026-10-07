import os
import io
import json
import csv
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field

@dataclass
class LoadedDocument:
    filename: str
    filepath: str
    file_type: str
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)
    sections: List[Dict[str, Any]] = field(default_factory=list)

class DocumentLoader:
    """
    Universal Multi-Format Document Parser with metadata and section extraction.
    Supports PDF, DOCX, Markdown, Plain Text, CSV, and JSON.
    """
    
    @staticmethod
    def load_file(file_path: str) -> LoadedDocument:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Document not found at path: {file_path}")
        
        filename = os.path.basename(file_path)
        ext = os.path.splitext(filename)[1].lower()
        
        if ext == ".pdf":
            return DocumentLoader._load_pdf(file_path, filename)
        elif ext in [".md", ".markdown"]:
            return DocumentLoader._load_markdown(file_path, filename)
        elif ext == ".txt":
            return DocumentLoader._load_txt(file_path, filename)
        elif ext == ".docx":
            return DocumentLoader._load_docx(file_path, filename)
        elif ext == ".csv":
            return DocumentLoader._load_csv(file_path, filename)
        elif ext == ".json":
            return DocumentLoader._load_json(file_path, filename)
        else:
            # Fallback to UTF-8 text reader
            return DocumentLoader._load_txt(file_path, filename)

    @staticmethod
    def load_from_bytes(file_bytes: bytes, filename: str) -> LoadedDocument:
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            return DocumentLoader._load_pdf_bytes(file_bytes, filename)
        elif ext in [".md", ".markdown", ".txt"]:
            text = file_bytes.decode("utf-8", errors="ignore")
            return LoadedDocument(
                filename=filename,
                filepath=f"memory://{filename}",
                file_type=ext.replace(".", ""),
                content=text,
                metadata={"size_bytes": len(file_bytes), "char_count": len(text)}
            )
        elif ext == ".csv":
            text = file_bytes.decode("utf-8", errors="ignore")
            return DocumentLoader._parse_csv_text(text, filename, f"memory://{filename}")
        elif ext == ".json":
            text = file_bytes.decode("utf-8", errors="ignore")
            return DocumentLoader._parse_json_text(text, filename, f"memory://{filename}")
        else:
            text = file_bytes.decode("utf-8", errors="ignore")
            return LoadedDocument(
                filename=filename,
                filepath=f"memory://{filename}",
                file_type="unknown",
                content=text,
                metadata={"size_bytes": len(file_bytes)}
            )

    @staticmethod
    def _load_pdf(file_path: str, filename: str) -> LoadedDocument:
        with open(file_path, "rb") as f:
            return DocumentLoader._load_pdf_bytes(f.read(), filename, file_path)

    @staticmethod
    def _load_pdf_bytes(file_bytes: bytes, filename: str, filepath: str = "") -> LoadedDocument:
        sections = []
        full_text = []
        num_pages = 0
        
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            num_pages = len(reader.pages)
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    full_text.append(page_text)
                    sections.append({
                        "section_id": f"page_{idx+1}",
                        "page_number": idx + 1,
                        "title": f"Page {idx+1}",
                        "text": page_text
                    })
        except Exception as e:
            # Simple text extraction fallback
            text = file_bytes.decode("latin-1", errors="ignore")
            import re
            cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\xff]', ' ', text)
            full_text = [cleaned]
            sections = [{"section_id": "fallback_1", "page_number": 1, "title": "Raw PDF Stream", "text": cleaned}]
            num_pages = 1

        content = "\n\n".join(full_text)
        return LoadedDocument(
            filename=filename,
            filepath=filepath or f"memory://{filename}",
            file_type="pdf",
            content=content,
            metadata={
                "num_pages": num_pages,
                "char_count": len(content),
                "word_count": len(content.split())
            },
            sections=sections
        )

    @staticmethod
    def _load_markdown(file_path: str, filename: str) -> LoadedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        
        # Header-aware section split
        sections = []
        lines = content.splitlines()
        current_section = {"title": "Introduction", "lines": []}
        
        for line in lines:
            if line.startswith("#"):
                if current_section["lines"]:
                    sec_text = "\n".join(current_section["lines"]).strip()
                    if sec_text:
                        sections.append({
                            "title": current_section["title"],
                            "text": sec_text
                        })
                current_section = {"title": line.lstrip("#").strip(), "lines": []}
            else:
                current_section["lines"].append(line)
        
        if current_section["lines"]:
            sec_text = "\n".join(current_section["lines"]).strip()
            if sec_text:
                sections.append({"title": current_section["title"], "text": sec_text})

        return LoadedDocument(
            filename=filename,
            filepath=file_path,
            file_type="markdown",
            content=content,
            metadata={
                "char_count": len(content),
                "word_count": len(content.split()),
                "section_count": len(sections)
            },
            sections=sections
        )

    @staticmethod
    def _load_txt(file_path: str, filename: str) -> LoadedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        return LoadedDocument(
            filename=filename,
            filepath=file_path,
            file_type="txt",
            content=content,
            metadata={"char_count": len(content), "word_count": len(content.split())}
        )

    @staticmethod
    def _load_docx(file_path: str, filename: str) -> LoadedDocument:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs)
        except Exception:
            with open(file_path, "rb") as f:
                content = f.read().decode("utf-8", errors="ignore")
        return LoadedDocument(
            filename=filename,
            filepath=file_path,
            file_type="docx",
            content=content,
            metadata={"char_count": len(content), "word_count": len(content.split())}
        )

    @staticmethod
    def _load_csv(file_path: str, filename: str) -> LoadedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return DocumentLoader._parse_csv_text(text, filename, file_path)

    @staticmethod
    def _parse_csv_text(text: str, filename: str, filepath: str) -> LoadedDocument:
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return LoadedDocument(filename=filename, filepath=filepath, file_type="csv", content="")
        
        headers = rows[0]
        formatted_rows = []
        for r_idx, row in enumerate(rows[1:]):
            row_dict = dict(zip(headers, row))
            row_str = " | ".join(f"{k}: {v}" for k, v in row_dict.items())
            formatted_rows.append(f"Row {r_idx+1}: {row_str}")
        
        content = f"Table Columns: {', '.join(headers)}\n\n" + "\n".join(formatted_rows)
        return LoadedDocument(
            filename=filename,
            filepath=filepath,
            file_type="csv",
            content=content,
            metadata={"num_rows": len(rows)-1, "num_cols": len(headers)}
        )

    @staticmethod
    def _load_json(file_path: str, filename: str) -> LoadedDocument:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return DocumentLoader._parse_json_text(text, filename, file_path)

    @staticmethod
    def _parse_json_text(text: str, filename: str, filepath: str) -> LoadedDocument:
        data = json.loads(text)
        formatted = json.dumps(data, indent=2)
        return LoadedDocument(
            filename=filename,
            filepath=filepath,
            file_type="json",
            content=formatted,
            metadata={"char_count": len(formatted)}
        )
